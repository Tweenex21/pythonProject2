from docx import Document
from django.http import HttpResponse
from GenerateDok.Dok.models import DocumentTemplate
from django.shortcuts import render


def home(request):
    return render(request, 'home.html')


def jinja_view(request):
    # Для Jinja2 шаблона
    return render(request, 'index.html', {'title': 'Пример Jinja2'})


def generate_docx(request, template_id):
    template = DocumentTemplate.objects.get(id=template_id)
    doc = Document(template.file.path)

    # Замена placeholders (пример для простого текста)
    for paragraph in doc.paragraphs:
        if '{{ client_name }}' in paragraph.text:
            paragraph.text = paragraph.text.replace('{{ client_name }}', 'Иванов Иван')
        if '{{ client_fio }}' in paragraph.text:
            paragraph.text = paragraph.text.replace('{{ client_fio }}', 'Степан')

    # Сохранение
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document')
    response['Content-Disposition'] = 'attachment; filename="generated.docx"'
    doc.save(response)
    return response

